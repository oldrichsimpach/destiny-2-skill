"""Read-only inspection of permitted XLSX exports. Does not recalculate formulas."""
import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.utils.cell import range_boundaries


def formula_text(cell):
    value = cell.value
    if cell.data_type == 'f':
        return value if isinstance(value, str) else getattr(value, 'text', str(value))
    return None


def inspect(path, selections=(), source_url=None, retrieved_at=None):
    path = Path(path)
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    try:
        result = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                  'source_url': source_url, 'retrieved_at': retrieved_at,
                  'inspected_at_utc': dt.datetime.now(dt.timezone.utc).isoformat(),
                  'recalculated': False, 'sheets': [], 'selections': [],
                  'named_ranges': {k: v.attr_text for k, v in formulas.defined_names.items()}}
        for sheet in formulas:
            count, blank, issues, errors = 0, 0, [], []
            for row in sheet:
                for cell in row:
                    expression = formula_text(cell)
                    value = cached[sheet.title][cell.coordinate].value
                    if expression:
                        count += 1
                        blank += value is None
                        if re.search(r'DUMMYFUNCTION|IMPORTRANGE|IMPORTXML|GOOGLEFINANCE|IMAGE\(', expression, re.I):
                            if len(issues) < 12:
                                issues.append({'cell': cell.coordinate, 'formula': expression})
                    if cached[sheet.title][cell.coordinate].data_type == 'e' and len(errors) < 12:
                        errors.append({'cell': cell.coordinate, 'error': value})
            result['sheets'].append({'name': sheet.title, 'state': sheet.sheet_state,
                'rows': sheet.max_row, 'columns': sheet.max_column, 'formula_count': count,
                'formula_caches_blank_or_missing': blank, 'nonportable_formula_examples': issues,
                'cached_error_examples': errors, 'protected': sheet.protection.sheet,
                'validations': [{'range': str(v.sqref), 'type': v.type, 'formula1': v.formula1,
                                 'formula2': v.formula2} for v in sheet.data_validations.dataValidation]})
        total_cells = 0
        for selection in selections:
            name, region = selection.rsplit('!', 1)
            name = name.strip("'").replace("''", "'")
            x1, y1, x2, y2 = range_boundaries(region)
            if None in (x1, y1, x2, y2) or x2 < x1 or y2 < y1:
                raise ValueError('Use a bounded rectangular cell range')
            total_cells += (x2-x1+1)*(y2-y1+1)
            if total_cells > 10000:
                raise ValueError('Requested ranges exceed 10000 cells')
            sheet = formulas[name]
            cells = []
            for row in sheet.iter_rows(min_row=y1, max_row=y2, min_col=x1, max_col=x2):
                for cell in row:
                    cells.append({'cell': cell.coordinate,
                        'value_or_formula': formula_text(cell) or cell.value,
                        'cached_value': cached[name][cell.coordinate].value,
                        'number_format': cell.number_format,
                        'comment': cell.comment.text if cell.comment else None})
            result['selections'].append({'sheet': name, 'range': region, 'cells': cells})
        return result
    finally:
        formulas.close()
        cached.close()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('path', type=Path)
    p.add_argument('--range', action='append', default=[], dest='ranges')
    p.add_argument('--source-url')
    p.add_argument('--retrieved-at', help='Actual acquisition timestamp; not inferred from file mtime')
    p.add_argument('--out', type=Path)
    a = p.parse_args()
    result = inspect(a.path, a.ranges, a.source_url, a.retrieved_at)
    text = json.dumps(result, ensure_ascii=False, indent=2, default=str)
    if a.out:
        a.out.write_text(text, encoding='utf-8')
    else:
        print(text)


if __name__ == '__main__':
    main()
